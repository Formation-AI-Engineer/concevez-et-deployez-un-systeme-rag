"""Génère le support de soutenance (étape 6.5) : ``docs/presentation.pptx``.

Le contenu est défini en clair ci-dessous (un dict par slide) puis rendu en PowerPoint via
``python-pptx``. Le support est ainsi **reproductible et versionnable** : on édite ce script
plutôt qu'un binaire, et on régénère le ``.pptx`` à la demande.

La charte graphique (couleurs navy/ambre, police Calibri, bandeau pied de page, kicker
« NN • SECTION ») reprend celle du support du Projet 8, pour une présentation homogène.

Usage :
    uv run --extra dev python scripts/build_presentation.py
    # → docs/presentation.pptx (ouvrable dans PowerPoint / Google Slides / LibreOffice)
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT_PATH = ROOT / "docs" / "presentation.pptx"

# --- Charte graphique (reprise du Projet 8) ---------------------------------------------------
NAVY = RGBColor(0x0F, 0x2F, 0x52)    # titres, fond de garde, bandeau pied de page
AMBER = RGBColor(0xF5, 0x9E, 0x0B)   # accents : kickers, barres, puces, chiffres clés
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF3, 0xF4, 0xF6)   # fond des encarts « chiffres clés »
LIGHT = RGBColor(0xE5, 0xE7, 0xEB)   # rail des barres de progression
DARK = RGBColor(0x1F, 0x29, 0x37)    # texte secondaire
BODY = RGBColor(0x37, 0x41, 0x51)    # texte des puces
FONT = "Calibri"

SLIDE_W = Inches(13.333)  # 16:9
SLIDE_H = Inches(7.5)
FOOTER_TEXT = "Projet 7 — Assistant RAG d'événements culturels  •  Lamine CAMARA"

# --- Contenu : un dict par slide --------------------------------------------------------------
# Slide de garde : kind="title". Slide standard : kicker + title + bullets (+ stats).
# Une puce est soit une str (niveau 0), soit un tuple (texte, niveau) pour un sous-niveau.
SLIDES: list[dict] = [
    {
        "kind": "title",
        "kicker": "PROJET 7  •  FORMATION AI ENGINEER  •  SYSTÈME RAG",
        "title": ["Assistant intelligent de", "recommandation d'événements culturels"],
        "title_size": 40,
        "subtitle": "Un système RAG pour Puls-Events",
        "tech": "LangChain • FAISS • Mistral • FastAPI • Docker",
        "author": ["Lamine CAMARA", "Soutenance — 2026"],
    },
    {
        "kicker": "DÉROULÉ",
        "title": "Déroulé de la présentation",
        "bullets": [
            "1 · Le système RAG développé — contexte, architecture, données, modèle.",
            "2 · Démonstration de l'API en direct (interrogation → réponse).",
            "3 · Résultats d'évaluation.",
            "4 · Structure du dépôt GitHub & principaux scripts.",
            "5 · Échanges & discussion.",
        ],
    },
    {
        "kicker": "01  •  CONTEXTE",
        "title": "Contexte & problème métier",
        "bullets": [
            "Puls-Events : plateforme de recommandation d'événements culturels.",
            "Besoin : un assistant qui répond en langage naturel sur les événements à venir.",
            "Ex. « Quels concerts de jazz à Paris ? », « Une expo près de la Cité des sciences ? ».",
            "Difficulté : une offre qui change sans cesse et des faits (dates, lieux) à ne pas se tromper.",
        ],
    },
    {
        "kicker": "02  •  PRINCIPE",
        "title": "Qu'est-ce qu'un système RAG  ?",
        "bullets": [
            "Un LLM seul répond « de mémoire » : risque d'événements faux ou périmés.",
            "RAG = Retrieval-Augmented Generation.",
            ("1) on RÉCUPÈRE d'abord des événements réels dans une base à jour ;", 1),
            ("2) le LLM RÉDIGE la réponse UNIQUEMENT à partir de ces événements.", 1),
            "Bénéfice : la fluidité d'un LLM + l'exactitude d'une base de données.",
        ],
    },
    {
        "kind": "objectives",
        "kicker": "03  •  OBJECTIFS",
        "title": "Objectifs du POC & périmètre",
        "objectives": [
            ("01", "Faisabilité technique",
             "De l'API Open Agenda à une réponse sourcée, une chaîne RAG complète de bout en bout."),
            ("02", "Valeur métier",
             "Des réponses pertinentes, honnêtes et directement exploitables par les équipes."),
            ("03", "Performance & reproductibilité",
             "Chargement unique des modèles, conteneurisation, environnement figé et rejouable."),
        ],
        "perimeter_label": "Périmètre du POC",
        "perimeter": [
            ("Paris", "zone ciblée"),
            ("< 1 an", "fenêtre temporelle"),
            ("~1500", "événements indexés"),
            ("4076", "chunks vectorisés"),
        ],
    },
    {
        "kind": "architecture",
        "kicker": "04  •  ARCHITECTURE",
        "title": "Architecture du système",
        "phase1_label": "INDEXATION   ·   hors-ligne",
        "row1": [
            ("Open Agenda", "API publique", False),
            ("Pré-traitement", "pandas → Parquet", False),
            ("Chunks + Embeddings", "MiniLM 384d · CPU", False),
            ("Index FAISS", "IndexFlatL2 · cosinus", True),
        ],
        "link": "L'index FAISS et le modèle d'embeddings sont chargés une seule fois au démarrage de l'API.",
        "phase2_label": "INTERROGATION   ·   temps réel",
        "row2": [
            ("Question", "langage naturel", False),
            ("Recherche FAISS", "top-k + seuil (gating)", True),
            ("Génération Mistral", "contexte → réponse", False),
            ("Réponse", "+ sources citées", False),
        ],
        "footnote": "Séparation nette : logique métier (rag/) vs exposition HTTP (api/).",
    },
    {
        "kicker": "05  •  DONNÉES",
        "title": "Des données brutes au jeu structuré",
        "bullets": [
            "Contrainte API : endpoint global /events fermé (403) → repli par agendas « Paris », puis re-filtrage sur la ville réelle des événements.",
            "Robustesse : pagination par curseur (bug d'offset contourné), retries + back-off, déduplication par identifiant.",
            "Nettoyage : champs multilingues (fr prioritaire), HTML décodé, dates normalisées, événements sans contenu écartés.",
            "Diversité : plafond par agenda pour limiter un biais thématique (~1/3 d'événements religieux).",
            "Sortie : 1500 événements (16 agendas), jeu typé au format Parquet — choisi pour son I/O rapide et typé.",
        ],
    },
    {
        "kicker": "06  •  VECTORISATION",
        "title": "Recherche par le sens",
        "bullets": [
            "Chunking : RecursiveCharacterTextSplitter (800 caractères, chevauchement 100).",
            "Embeddings : sentence-transformers multilingue (MiniLM-L12-v2, 384 dim), en LOCAL sur CPU.",
            "Vecteurs normalisés → distance L2 ≈ similarité cosinus.",
            "Index FAISS IndexFlatL2 : recherche exacte, rappel 100 %, instantanée à cette échelle.",
            "Métadonnées stockées avec chaque vecteur (titre, date, lieu, url) → restitution des sources.",
        ],
    },
    {
        "kicker": "07  •  GÉNÉRATION",
        "title": "Mistral + garde-fous anti-hallucination",
        "bullets": [
            "LLM : Mistral (mistral-small-latest) — qualité en français, coût maîtrisé, intégré à LangChain.",
            "Prompt système : répondre EXCLUSIVEMENT à partir du contexte, ne rien inventer, citer date/lieu.",
            "Gating : si aucun événement pertinent (seuil cosinus) → refus honnête SANS appeler le LLM.",
            "Résultat : pas d'invention, et une réponse déterministe + économique hors-périmètre.",
        ],
    },
    {
        "kind": "api",
        "kicker": "08  •  API",
        "title": "API REST & robustesse",
        "endpoints": [
            ("GET", "/health", "état du service + nb d'événements indexés"),
            ("POST", "/ask", "question → réponse + sources citées"),
            ("POST", "/rebuild", "reconstruit l'index (jeton X-API-Token)"),
            ("GET", "/docs", "documentation Swagger auto-générée"),
        ],
        "single_load": "Index FAISS + modèle d'embeddings chargés une seule fois au démarrage (lifespan).",
        "errors": [
            ("422", "question vide / champ manquant"),
            ("503", "index absent (assistant non chargé)"),
            ("401", "jeton /rebuild invalide"),
            ("500", "erreur interne — détail jamais exposé"),
        ],
        "quality": [
            "82 tests automatisés (unitaires + API).",
            "CI GitHub Actions : lint + tests à chaque push.",
            "Jeton /rebuild comparé en temps constant.",
        ],
    },
    {
        "kicker": "09  •  DÉMO",
        "title": "Démonstration live de l'API",
        "bullets": [
            "Scénario 1 : « Quels concerts de jazz à Paris ? » → Jazzycolors, Café Maa…",
            "Scénario 2 : « Une expo de peinture à voir à Paris ? » → Mai-Thu Perret…",
            "Scénario 3 : question hors-périmètre → refus honnête (pas d'invention).",
            "Le tout servi par le conteneur Docker, en local (Swagger /docs).",
        ],
    },
    {
        "kind": "results",
        "kicker": "10  •  RÉSULTATS",
        "title": "Résultats d'évaluation",
        "metrics": [
            ("0.77", "Similarité\nsémantique moy."),
            ("0.71", "Couverture\ndes infos clés"),
            ("14/20", "Réponses\ncorrectes"),
            ("20", "Questions\nannotées"),
        ],
        "bars_title": "Réussite par catégorie",
        "bars": [
            ("Lieu (n=5)", 0.8, "80 %", False),
            ("Période (n=2)", 0.5, "50 %", False),
            ("Type d'événement", 0.8, "80 %", False),
            ("Hors-périmètre", 0.33, "33 %*", True),
        ],
        "aside_title": "Recul critique",
        "aside": [
            "* Un refus hors-périmètre — pourtant le bon comportement — est compté « incorrect » par la métrique.",
            "Les 14/20 sont donc un plancher pessimiste ; détail complet dans le rapport technique.",
        ],
    },
    {
        "kind": "repo",
        "kicker": "11  •  DÉPÔT",
        "title": "Structure du dépôt & principaux scripts",
        "tree": (
            "projet7/\n"
            "├── rag/         cœur métier RAG\n"
            "│      data_loader · preprocessing\n"
            "│      chunking · embeddings\n"
            "│      vectorstore · chain · config\n"
            "├── api/         FastAPI (main, schemas)\n"
            "├── scripts/     pipeline & outils CLI\n"
            "├── tests/       82 tests (10 fichiers)\n"
            "├── eval/        jeu annoté + rapports\n"
            "├── docs/        fiches + rapport tech.\n"
            "├── Dockerfile · docker-compose.yml\n"
            "└── pyproject.toml · uv.lock"
        ),
        "pipeline": [
            "fetch_events.py",
            "preprocess_events.py",
            "build_index.py",
        ],
        "tools": [
            "search.py — recherche sémantique en CLI.",
            "evaluate_rag.py — évaluation de la qualité.",
            "check_imports.py — vérification des dépendances.",
            "build_presentation.py — ce support.",
        ],
    },
    {
        "kicker": "12  •  DÉPLOIEMENT",
        "title": "Conteneurisation & reproductibilité",
        "bullets": [
            "Image Docker : API + index FAISS + modèle d'embeddings EMBARQUÉS → démarrage hors-ligne.",
            "Seule la génération Mistral appelle le réseau, au moment de la requête.",
            "PyTorch CPU épinglé (projet 100 % CPU) → image ~3,8 Go au lieu de ~11,6 Go.",
            "Environnement figé avec uv (pyproject + uv.lock) → installs déterministes.",
            "Secrets jamais versionnés ni copiés dans l'image : injectés au run via --env-file.",
        ],
    },
    {
        "kind": "roadmap",
        "kicker": "13  •  PERSPECTIVES",
        "title": "Limites & pistes d'industrialisation",
        "left_header": "LIMITE ACTUELLE",
        "right_header": "PISTE D'INDUSTRIALISATION",
        "rows": [
            ("Couverture : 1 ville, instantané figé des données.",
             "Fraîcheur : ingestion planifiée + /rebuild automatisé."),
            ("Rappel : requêtes par type d'événement parfois incomplètes.",
             "Recherche hybride (lexicale + vectorielle) + reranking."),
            ("Coût & latence de l'appel au LLM externe (Mistral).",
             "Cache de réponses, quotas, observabilité des coûts LLM."),
            ("Recherche exacte FAISS dimensionnée pour le POC (~4K vecteurs).",
             "Index IVF/HNSW, base vectorielle managée, auth & rate-limiting."),
        ],
    },
    {
        "kind": "title",
        "kicker": "DISCUSSION  •  MERCI",
        "title": ["Questions  ?"],
        "subtitle": "Merci de votre attention — échanges & discussion.",
        "tech": "Code · rapport technique · fiches d'étape disponibles dans le dépôt",
        "author": ["Lamine CAMARA", "Assistant RAG — Puls-Events (POC)"],
    },
]


# --- Primitives de rendu ----------------------------------------------------------------------
def _rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _text(slide, left, top, width, height, runs, *, size, color, bold=False,
          anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, space_after=0, line=None, font=FONT):
    """`runs` = liste de paragraphes ; chaque paragraphe = str ou liste de (texte, couleur, bold)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line:
            p.line_spacing = line
        segments = [(para, color, bold)] if isinstance(para, str) else para
        for txt, col, bd in segments:
            r = p.add_run()
            r.text = txt
            r.font.name = font
            r.font.size = Pt(size)
            r.font.bold = bd
            r.font.color.rgb = col
    return box


def _footer(slide, page, total):
    _rect(slide, 0, Inches(7.15), SLIDE_W, Inches(0.35), NAVY)
    _text(slide, Inches(0.4), Inches(7.17), Inches(10.5), Inches(0.3),
          [FOOTER_TEXT], size=10, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    _text(slide, Inches(12.1), Inches(7.17), Inches(1.0), Inches(0.3),
          [f"{page} / {total}"], size=10, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


def _stats_panel(slide, stats):
    """Encart « chiffres clés » à droite : grand nombre ambre + libellé."""
    _rect(slide, Inches(8.55), Inches(2.0), Inches(4.1), Inches(4.7), PANEL)
    _text(slide, Inches(8.85), Inches(2.15), Inches(3.6), Inches(0.5),
          ["Chiffres clés"], size=16, color=NAVY, bold=True)
    y = 2.85
    for value, label in stats:
        _text(slide, Inches(8.85), Inches(y), Inches(1.7), Inches(0.6),
              [value], size=24, color=AMBER, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _text(slide, Inches(10.5), Inches(y), Inches(2.0), Inches(0.6),
              [label], size=12, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
        y += 0.92


def _title_slide(slide, s):
    _rect(slide, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    _rect(slide, 0, Inches(6.7), SLIDE_W, Inches(0.15), AMBER)
    _text(slide, Inches(0.9), Inches(1.3), Inches(11.5), Inches(0.4),
          [s["kicker"]], size=14, color=AMBER, bold=True)
    _text(slide, Inches(0.9), Inches(2.0), Inches(11.7), Inches(2.2),
          s["title"], size=s.get("title_size", 50), color=WHITE, bold=True, line=1.05, space_after=2)
    _text(slide, Inches(0.9), Inches(4.4), Inches(11.0), Inches(0.6),
          [s["subtitle"]], size=24, color=WHITE)
    _text(slide, Inches(0.9), Inches(5.05), Inches(11.5), Inches(0.4),
          [s["tech"]], size=16, color=AMBER, bold=True)
    _text(slide, Inches(0.9), Inches(5.9), Inches(11.0), Inches(0.9),
          s["author"], size=14, color=WHITE, space_after=2)


def _slide_header(slide, s):
    """Bandeau commun : barre latérale ambre, kicker, titre, soulignement."""
    _rect(slide, 0, 0, Inches(0.3), SLIDE_H, AMBER)  # barre latérale
    _text(slide, Inches(0.7), Inches(0.55), Inches(12.0), Inches(0.3),
          [s["kicker"]], size=11, color=AMBER, bold=True)
    _text(slide, Inches(0.7), Inches(0.9), Inches(12.0), Inches(0.7),
          [s["title"]], size=32, color=NAVY, bold=True)
    _rect(slide, Inches(0.72), Inches(1.7), Inches(1.2), Inches(0.04), AMBER)  # soulignement


def _metric_cards(slide, metrics, top=2.05):
    """Bandeau horizontal de cartes « tuiles » navy : grand nombre ambre + libellé blanc.

    Reprend le vocabulaire visuel des encarts du Projet 8 (slide architecture)."""
    card_w, card_h, gap = 2.75, 1.62, 0.33
    start = 0.85
    for i, (value, label) in enumerate(metrics):
        left = start + i * (card_w + gap)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(card_w), Inches(card_h))
        card.fill.solid()
        card.fill.fore_color.rgb = NAVY
        card.line.fill.background()
        card.shadow.inherit = False
        _rect(slide, Inches(left), Inches(top), Inches(card_w), Inches(0.08), AMBER)  # liseré haut
        _text(slide, Inches(left), Inches(top + 0.22), Inches(card_w), Inches(0.78),
              [value], size=34, color=AMBER, bold=True,
              anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        _text(slide, Inches(left + 0.1), Inches(top + 1.0), Inches(card_w - 0.2), Inches(0.55),
              label.split("\n"), size=12, color=WHITE,
              anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.CENTER, line=1.0)


def _roadmap_slide(slide, s, page, total):
    """Slide perspectives : mapping « Limite actuelle → Piste d'industrialisation »."""
    _slide_header(slide, s)
    left_x, left_w = 0.85, 5.05
    right_x, right_w = 6.7, 6.05
    _text(slide, Inches(left_x + 0.05), Inches(2.05), Inches(left_w), Inches(0.32),
          [s["left_header"]], size=12, color=DARK, bold=True)
    _text(slide, Inches(right_x + 0.05), Inches(2.05), Inches(right_w), Inches(0.32),
          [s["right_header"]], size=12, color=AMBER, bold=True)

    row_h, gap, y = 0.92, 0.13, 2.5
    for limit, fix in s["rows"]:
        for x, w, txt, fill, fg in (
            (left_x, left_w, limit, PANEL, NAVY),
            (right_x, right_w, fix, NAVY, WHITE),
        ):
            box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(row_h))
            box.fill.solid()
            box.fill.fore_color.rgb = fill
            box.line.fill.background()
            box.shadow.inherit = False
            _text(slide, Inches(x + 0.25), Inches(y), Inches(w - 0.5), Inches(row_h),
                  [txt], size=12.5, color=fg, anchor=MSO_ANCHOR.MIDDLE, line=1.05)
        arr = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW, Inches(left_x + left_w + 0.08),
            Inches(y + row_h / 2 - 0.13), Inches(right_x - left_x - left_w - 0.16), Inches(0.26))
        arr.fill.solid()
        arr.fill.fore_color.rgb = AMBER
        arr.line.fill.background()
        arr.shadow.inherit = False
        y += row_h + gap

    _footer(slide, page, total)


def _objective_card(slide, left, top, w, h, num, title, body):
    """Carte d'objectif : panneau clair, liseré ambre, badge numéroté, titre + description."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = PANEL
    card.line.fill.background()
    card.shadow.inherit = False
    _rect(slide, Inches(left), Inches(top), Inches(w), Inches(0.08), AMBER)  # liseré haut
    _text(slide, Inches(left + 0.3), Inches(top + 0.22), Inches(w - 0.6), Inches(0.45),
          [num], size=22, color=AMBER, bold=True)
    _text(slide, Inches(left + 0.3), Inches(top + 0.72), Inches(w - 0.55), Inches(0.65),
          [title], size=15, color=NAVY, bold=True, line=1.0)
    _text(slide, Inches(left + 0.3), Inches(top + 1.42), Inches(w - 0.55), Inches(1.0),
          [body], size=12.5, color=BODY, line=1.12)


def _objectives_slide(slide, s, page, total):
    """Slide objectifs : 3 cartes d'objectifs + bande de tuiles « périmètre »."""
    _slide_header(slide, s)
    card_w, gap, start, top, h = 3.73, 0.35, 0.75, 2.05, 2.5
    for i, (num, title, body) in enumerate(s["objectives"]):
        _objective_card(slide, start + i * (card_w + gap), top, card_w, h, num, title, body)

    _text(slide, Inches(0.85), Inches(4.72), Inches(11.9), Inches(0.35),
          [s.get("perimeter_label", "Périmètre du POC")], size=15, color=NAVY, bold=True)
    _metric_cards(slide, s["perimeter"], top=5.2)

    _footer(slide, page, total)


def _results_slide(slide, s, page, total):
    """Slide « résultats » soignée : cartes de métriques + lecture + encart de recul critique."""
    _slide_header(slide, s)
    _metric_cards(slide, s["metrics"])

    # Colonne gauche : réussite par catégorie (mini-barres de progression).
    _text(slide, Inches(0.85), Inches(3.95), Inches(7.0), Inches(0.4),
          [s["bars_title"]], size=16, color=NAVY, bold=True)
    bar_x, bar_w, y = 3.45, 2.9, 4.6
    for label, pct, value, highlight in s["bars"]:
        _text(slide, Inches(0.85), Inches(y - 0.04), Inches(2.45), Inches(0.32),
              [label], size=12.5, color=BODY, anchor=MSO_ANCHOR.MIDDLE)
        _rect(slide, Inches(bar_x), Inches(y), Inches(bar_w), Inches(0.26), LIGHT)
        _rect(slide, Inches(bar_x), Inches(y), Inches(bar_w * pct), Inches(0.26),
              AMBER if highlight else NAVY)
        _text(slide, Inches(bar_x + bar_w + 0.15), Inches(y - 0.04), Inches(0.95), Inches(0.32),
              [value], size=12.5, color=AMBER if highlight else NAVY, bold=True,
              anchor=MSO_ANCHOR.MIDDLE)
        y += 0.52

    # Colonne droite : encart « recul critique ».
    panel_l, panel_t, panel_w, panel_h = 8.15, 3.95, 4.5, 2.78
    _rect(slide, Inches(panel_l), Inches(panel_t), Inches(panel_w), Inches(panel_h), PANEL)
    _rect(slide, Inches(panel_l), Inches(panel_t), Inches(0.1), Inches(panel_h), AMBER)  # accent
    _text(slide, Inches(panel_l + 0.3), Inches(panel_t + 0.2), Inches(panel_w - 0.5), Inches(0.4),
          [s["aside_title"]], size=15, color=NAVY, bold=True)
    _text(slide, Inches(panel_l + 0.3), Inches(panel_t + 0.75), Inches(panel_w - 0.55),
          Inches(panel_h - 0.9), s["aside"], size=12.5, color=DARK, space_after=8, line=1.12)

    _footer(slide, page, total)


def _flow_box(slide, left, top, w, h, title, sub, highlight):
    """Boîte d'étape du diagramme de flux (navy, ou ambre si maillon mis en avant)."""
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = AMBER if highlight else NAVY
    box.line.fill.background()
    box.shadow.inherit = False
    title_c = NAVY if highlight else WHITE
    sub_c = DARK if highlight else PANEL
    _text(slide, Inches(left), Inches(top + 0.2), Inches(w), Inches(0.52),
          [title], size=14, color=title_c, bold=True,
          anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    _text(slide, Inches(left + 0.08), Inches(top + 0.72), Inches(w - 0.16), Inches(0.4),
          [sub], size=10.5, color=sub_c, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.CENTER)


def _flow_row(slide, items, top):
    """Rangée de boîtes reliées par des flèches ambre."""
    box_w, box_h, gap, start = 2.6, 1.2, 0.5, 0.75
    for i, (title, sub, highlight) in enumerate(items):
        left = start + i * (box_w + gap)
        _flow_box(slide, left, top, box_w, box_h, title, sub, highlight)
        if i < len(items) - 1:
            arr = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, Inches(left + box_w + 0.06),
                Inches(top + box_h / 2 - 0.13), Inches(gap - 0.12), Inches(0.26))
            arr.fill.solid()
            arr.fill.fore_color.rgb = AMBER
            arr.line.fill.background()
            arr.shadow.inherit = False


def _architecture_slide(slide, s, page, total):
    """Slide architecture en diagramme de flux : deux pipelines (indexation / interrogation)."""
    _slide_header(slide, s)
    caption = lambda y, txt: _text(  # noqa: E731
        slide, Inches(0.75), Inches(y), Inches(11.9), Inches(0.4),
        [[("▸  ", AMBER, True), (txt, BODY, False)]], size=12.5, color=BODY)

    _text(slide, Inches(0.75), Inches(2.05), Inches(11.9), Inches(0.3),
          [s["phase1_label"]], size=12, color=AMBER, bold=True)
    _flow_row(slide, s["row1"], 2.5)
    caption(3.95, s["link"])
    _text(slide, Inches(0.75), Inches(4.55), Inches(11.9), Inches(0.3),
          [s["phase2_label"]], size=12, color=AMBER, bold=True)
    _flow_row(slide, s["row2"], 5.0)
    caption(6.45, s["footnote"])

    _footer(slide, page, total)


def _api_slide(slide, s, page, total):
    """Slide API : table d'endpoints (badges + chemins monospace) + panneaux erreurs / qualité."""
    _slide_header(slide, s)

    # Colonne gauche : endpoints.
    _text(slide, Inches(0.85), Inches(2.05), Inches(6.3), Inches(0.35),
          ["Endpoints"], size=15, color=NAVY, bold=True)
    y = 2.62
    for method, path, desc in s["endpoints"]:
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(y), Inches(0.92), Inches(0.44))
        badge.fill.solid()
        badge.fill.fore_color.rgb = AMBER if method == "POST" else NAVY
        badge.line.fill.background()
        badge.shadow.inherit = False
        _text(slide, Inches(0.85), Inches(y), Inches(0.92), Inches(0.44),
              [method], size=11, color=WHITE, bold=True,
              anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        _text(slide, Inches(1.92), Inches(y), Inches(1.95), Inches(0.44),
              [path], size=14, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE, font="Consolas")
        _text(slide, Inches(3.95), Inches(y), Inches(3.3), Inches(0.44),
              [desc], size=11.5, color=BODY, anchor=MSO_ANCHOR.MIDDLE)
        y += 0.72
    _text(slide, Inches(0.85), Inches(y + 0.05), Inches(6.4), Inches(0.6),
          [[("▸  ", AMBER, True), (s["single_load"], BODY, False)]], size=11.5, color=BODY, line=1.05)

    # Colonne droite : erreurs (haut) + qualité (bas).
    px, pw = 7.55, 5.25
    _rect(slide, Inches(px), Inches(2.05), Inches(pw), Inches(2.55), PANEL)
    _rect(slide, Inches(px), Inches(2.05), Inches(0.1), Inches(2.55), AMBER)
    _text(slide, Inches(px + 0.3), Inches(2.2), Inches(pw - 0.5), Inches(0.35),
          ["Gestion des erreurs"], size=14, color=NAVY, bold=True)
    ey = 2.78
    for code, meaning in s["errors"]:
        cb = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(px + 0.3), Inches(ey), Inches(0.72), Inches(0.34))
        cb.fill.solid()
        cb.fill.fore_color.rgb = NAVY
        cb.line.fill.background()
        cb.shadow.inherit = False
        _text(slide, Inches(px + 0.3), Inches(ey), Inches(0.72), Inches(0.34),
              [code], size=11, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        _text(slide, Inches(px + 1.15), Inches(ey), Inches(pw - 1.4), Inches(0.34),
              [meaning], size=11.5, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
        ey += 0.45

    _rect(slide, Inches(px), Inches(4.8), Inches(pw), Inches(1.95), PANEL)
    _rect(slide, Inches(px), Inches(4.8), Inches(0.1), Inches(1.95), AMBER)
    _text(slide, Inches(px + 0.3), Inches(4.95), Inches(pw - 0.5), Inches(0.35),
          ["Qualité & robustesse"], size=14, color=NAVY, bold=True)
    qp = [[("▸  ", AMBER, True), (t, DARK, False)] for t in s["quality"]]
    _text(slide, Inches(px + 0.3), Inches(5.45), Inches(pw - 0.55), Inches(1.2),
          qp, size=11.5, color=DARK, space_after=6, line=1.08)

    _footer(slide, page, total)


def _repo_slide(slide, s, page, total):
    """Slide dépôt : arborescence monospace (gauche) + pipeline de scripts (droite)."""
    _slide_header(slide, s)

    # Colonne gauche : arborescence.
    _rect(slide, Inches(0.85), Inches(2.05), Inches(6.35), Inches(4.7), PANEL)
    _rect(slide, Inches(0.85), Inches(2.05), Inches(0.1), Inches(4.7), AMBER)
    _text(slide, Inches(1.1), Inches(2.2), Inches(6.0), Inches(4.4),
          s["tree"].split("\n"), size=12, color=DARK, font="Consolas", line=1.18)

    # Colonne droite : pipeline d'indexation (flux vertical) + autres scripts.
    px, pw = 7.6, 5.2
    _text(slide, Inches(px), Inches(2.05), Inches(pw), Inches(0.35),
          ["Pipeline d'indexation"], size=15, color=NAVY, bold=True)
    y = 2.6
    for i, step in enumerate(s["pipeline"]):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(px), Inches(y), Inches(pw), Inches(0.55))
        box.fill.solid()
        box.fill.fore_color.rgb = NAVY
        box.line.fill.background()
        box.shadow.inherit = False
        _text(slide, Inches(px + 0.25), Inches(y), Inches(pw - 0.4), Inches(0.55),
              [[(f"{i + 1} · ", AMBER, True), (step, WHITE, False)]], size=12.5, color=WHITE,
              anchor=MSO_ANCHOR.MIDDLE, font="Consolas")
        if i < len(s["pipeline"]) - 1:
            arr = slide.shapes.add_shape(
                MSO_SHAPE.DOWN_ARROW, Inches(px + pw / 2 - 0.11), Inches(y + 0.57),
                Inches(0.22), Inches(0.16))
            arr.fill.solid()
            arr.fill.fore_color.rgb = AMBER
            arr.line.fill.background()
            arr.shadow.inherit = False
        y += 0.73
    _text(slide, Inches(px), Inches(y + 0.05), Inches(pw), Inches(0.35),
          ["Autres scripts"], size=15, color=NAVY, bold=True)
    op = [[("▸  ", AMBER, True), (t, BODY, False)] for t in s["tools"]]
    _text(slide, Inches(px), Inches(y + 0.5), Inches(pw), Inches(1.4),
          op, size=12, color=BODY, space_after=6, line=1.08)

    _footer(slide, page, total)


def _content_slide(slide, s, page, total):
    _slide_header(slide, s)

    has_stats = bool(s.get("stats"))
    body_w = Inches(7.4) if has_stats else Inches(11.9)
    paras = []
    for item in s["bullets"]:
        text, level = item if isinstance(item, tuple) else (item, 0)
        marker = "▸  " if level == 0 else "–  "
        paras.append([(marker, AMBER, True), (text, BODY, False)])
    _text(slide, Inches(0.85), Inches(2.15), body_w, Inches(4.6),
          paras, size=17, color=BODY, space_after=12, line=1.1)

    if has_stats:
        _stats_panel(slide, s["stats"])
    _footer(slide, page, total)



def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    total = len(SLIDES)

    for i, s in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        if s.get("kind") == "title":
            _title_slide(slide, s)
        elif s.get("kind") == "objectives":
            _objectives_slide(slide, s, i, total)
        elif s.get("kind") == "architecture":
            _architecture_slide(slide, s, i, total)
        elif s.get("kind") == "results":
            _results_slide(slide, s, i, total)
        elif s.get("kind") == "roadmap":
            _roadmap_slide(slide, s, i, total)
        elif s.get("kind") == "api":
            _api_slide(slide, s, i, total)
        elif s.get("kind") == "repo":
            _repo_slide(slide, s, i, total)
        else:
            _content_slide(slide, s, i, total)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PATH))
    return OUT_PATH


if __name__ == "__main__":
    path = build()
    print(f"Présentation générée : {path}  ({len(SLIDES)} slides)")
